import json
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from studybr.importers import build_diagnostic, serialize_dataset
from studybr.importers.diagnostic import normalized_path
from studybr.importers.workbooks import CREATION_TABLES, SourceError, read_workbook
from studybr.models import PageKind
from tests.fixtures.factories import source_pair

ROOT = Path(__file__).resolve().parents[2]
DOMAIN = ROOT / "config/domain.json"


def build(pair, **kwargs):
    return build_diagnostic(
        *pair, base_url="https://example.test/studyinbrazil", domain_path=DOMAIN, **kwargs
    )


def test_explicit_tables_unicode_nulls_quotes_and_determinism(tmp_path):
    pair = source_pair(tmp_path, name=unicodedata.normalize("NFD", 'Universidade "São José"'))
    first = build(pair)
    second = build(pair)
    assert serialize_dataset(first) == serialize_dataset(second)
    assert len(first.institutions) == 1
    assert first.institutions[0].official_name == 'Universidade "São José"'
    assert len(first.institutions[0].localizations) == 3
    assert sum(len(loc.pages) for loc in first.institutions[0].localizations) == 12
    assert all(
        page.description is None
        for loc in first.institutions[0].localizations
        for page in loc.pages
    )
    assert "São José" in serialize_dataset(first)


def test_editorial_relation_is_never_silently_international_students(tmp_path):
    dataset = build(source_pair(tmp_path))
    assert len(dataset.unassigned_editorial_rows) == 1
    assert any(c.category == "editorial" for c in dataset.conflicts)
    for loc in dataset.institutions[0].localizations:
        page = next(p for p in loc.pages if p.semantic_id == PageKind.INTERNATIONAL_STUDENTS)
        assert page.editorial_policy.phase is None


def test_source_name_conflict_keeps_all_candidates(tmp_path):
    pair = source_pair(tmp_path)
    workbook = load_workbook(pair[1])
    workbook["siglas,nomes pt en es"]["B2"] = "Nome alternativo"
    workbook.save(pair[1])
    dataset = build(pair)
    assert dataset.institutions[0].official_name is None
    assert any(c.category == "names" and c.field == "official_name" for c in dataset.conflicts)
    assert "Nome alternativo" in serialize_dataset(dataset)


def test_aliases_are_recorded_not_used_to_derive_slug(tmp_path):
    dataset = build(source_pair(tmp_path, acronym="CEFET_MG"))
    institution = dataset.institutions[0]
    assert institution.acronym == "CEFET-MG"
    assert "CEFET_MG" in institution.aliases
    assert any(c.category == "aliases" for c in dataset.conflicts)
    assert institution.localizations[0].home_slug == "ufscar"


def test_missing_named_sheet_is_error_not_first_sheet_fallback(tmp_path):
    pair = source_pair(tmp_path)
    workbook = load_workbook(pair[0])
    workbook["Português - home"].title = "Other"
    workbook.save(pair[0])
    with pytest.raises(SourceError, match="worksheet missing"):
        build(pair)


@pytest.mark.parametrize("cache", [None, "ufscar"])
def test_formula_cache_is_never_assumed_certified(tmp_path, cache):
    pair = source_pair(tmp_path)
    workbook = load_workbook(pair[0])
    # B is Sigla, K is the slug in this synthetic home table.
    workbook["Português - home"]["K3"] = "=LOWER(B3)"
    workbook.save(pair[0])
    if cache is not None:
        with zipfile.ZipFile(pair[0]) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        tree = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        cell = tree.find('.//s:c[@r="K3"]', ns)
        cell.set("t", "str")
        cell.find("s:v", ns).text = cache
        entries["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
        with zipfile.ZipFile(pair[0], "w") as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
    _, conflicts, _ = read_workbook(pair[0], "creation", CREATION_TABLES)
    conflict = next(c for c in conflicts if c.category == "formula_cache")
    assert conflict.candidates[0].provenance.cache_state == (
        "missing" if cache is None else "unverified"
    )


def test_duplicate_destination_is_not_overwritten(tmp_path):
    pair = source_pair(tmp_path)
    workbook = load_workbook(pair[0])
    sheet = workbook["Português - home"]
    values = [sheet.cell(3, col).value for col in range(1, sheet.max_column + 1)]
    values[1] = "OUTRA"
    sheet.append(values)
    table = sheet.tables["Tabela1"]
    table.ref = f"B2:{sheet.cell(4, sheet.max_column).coordinate}"
    workbook.save(pair[0])
    dataset = build(pair)
    assert any(c.category == "duplicate_destination" for c in dataset.conflicts)
    assert len(dataset.institutions) == 2


def test_csv_slug_and_metadata_divergences_are_preserved(tmp_path):
    pair = source_pair(tmp_path)
    csv_path = tmp_path / "pages.csv"
    csv_path.write_text(
        "Sigla,tipo3,nome_curto,Onde criar,Título,Descrição\nUFSCar,sobre_nos,sobre-nos,/pt-br/institutions/state/ufscar,Título novo,Descrição nova\n",
        encoding="utf-8",
    )
    dataset = build(pair, csv_sources={"subpage": csv_path})
    page = next(
        p
        for loc in dataset.institutions[0].localizations
        if loc.language == "pt-br"
        for p in loc.pages
        if p.semantic_id == "about"
    )
    assert page.slug is None
    assert page.title is None
    assert any(c.category == "xlsx_csv" for c in dataset.conflicts)
    assert any(c.category == "slugs" for c in dataset.conflicts)


def test_equivalent_path_forms_normalize_without_duplicating_site_prefix():
    base = "https://example.test/studyinbrazil"
    paths = [
        "https://example.test/studyinbrazil/en/institutions",
        "/studyinbrazil/en/institutions",
        "/en/institutions",
    ]
    assert {normalized_path(p, base) for p in paths} == {"/en/institutions"}


def test_external_source_path_is_conflict_not_selected(tmp_path):
    pair = source_pair(tmp_path)
    workbook = load_workbook(pair[0])
    workbook["Português - home"]["L3"] = "https://external.example/site"
    workbook.save(pair[0])
    dataset = build(pair)
    assert any(c.category == "paths" for c in dataset.conflicts)


def test_output_redacts_input_secrets(tmp_path):
    pair = source_pair(tmp_path, name="Bearer sensitive-input")
    serialized = serialize_dataset(build(pair))
    assert "sensitive-input" not in serialized
    assert "[REDACTED]" in serialized
    json.loads(serialized)
