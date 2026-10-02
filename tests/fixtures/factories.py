"""Small synthetic workbooks using exactly the supported table names."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from studybr.importers.workbooks import CREATION_TABLES, PRODUCTION_TABLES


def table(workbook, spec, heads, rows):
    sheet = workbook.create_sheet(spec.sheet)
    if spec.role in {"home", "subpage"}:
        sheet.append([None])
        # Real creation tables begin in B2.
        sheet.append([None] + heads)
        for row in rows:
            sheet.append([None] + row)
        ref = f"B2:{sheet.cell(sheet.max_row, len(heads) + 1).coordinate}"
    else:
        sheet.append(heads)
        for row in rows:
            sheet.append(row)
        ref = f"A1:{sheet.cell(sheet.max_row, len(heads)).coordinate}"
    sheet.add_table(Table(displayName=spec.table, ref=ref))


def source_pair(tmp_path: Path, *, acronym="UFSCar", name='Universidade "São José"'):
    creation = Workbook()
    creation.remove(creation.active)
    home_specs = {}
    for spec in CREATION_TABLES:
        if spec.role == "home":
            heads = [
                "Sigla",
                "Tipo",
                "Tipo2",
                "Natureza",
                "Universidade",
                "Descrição",
                "Região",
                "Estado",
                "Título",
                "nome_curto",
                "Onde criar",
                "TAGs",
            ]
            values = [
                acronym,
                "Universidade",
                "Federal",
                "Pública",
                name,
                None,
                "Sudeste",
                "SP",
                f"{acronym} - {name}",
                "ufscar",
                f"/{spec.language}/institutions/state",
                f"{acronym}; Universidade",
            ]
            if spec.language != "pt-br":
                heads += ["Nome Inglês" if spec.language == "en" else "Nome Espanhol"]
                values += [name]
            home_specs[spec.language] = dict(zip(heads, values, strict=True))
            table(creation, spec, heads, [values])
        elif spec.role == "subpage":
            heads = [
                "Sigla",
                "tipo3" if spec.language != "es" else "Tipo3",
                "nome_curto",
                "Onde criar",
                "Título" if spec.language != "es" else "Título2",
                "TAGs",
            ]
            identifiers = {
                "pt-br": ["sobre_nos", "vida_na_ies", "alunos_internacionais"],
                "en": ["about_us", "life_at_the_hei", "international_students"],
                "es": ["quienes_somos", "vida_en_la_ies", "estudiantes_internacionales"],
            }[spec.language]
            values = [
                [
                    acronym,
                    slug,
                    slug,
                    f"/{spec.language}/institutions/state/ufscar",
                    f"Página {slug}",
                    acronym,
                ]
                for slug in identifiers
            ]
            table(creation, spec, heads, values)
        else:
            heads = list(spec.required) + ["TAGs"]
            # Compiled rows copy the primary destination to test matching and provenance.
            if spec.role == "compiled_home":
                data = home_specs["pt-br"]
                values = [
                    [
                        "Home Portugues",
                        data["Onde criar"],
                        data["Título"],
                        data["nome_curto"],
                        data["TAGs"],
                    ]
                ]
            else:
                values = [
                    [
                        "subpáginas português",
                        "/pt-br/institutions/state/ufscar",
                        "Página sobre_nos",
                        "sobre_nos",
                        acronym,
                    ]
                ]
            table(creation, spec, heads, values)
    production = Workbook()
    production.remove(production.active)
    for spec in PRODUCTION_TABLES:
        if spec.role == "names":
            table(production, spec, list(spec.required), [[acronym.upper(), name, name, name]])
        elif spec.role == "editorial":
            heads = ["IES", "Página", "PT - Fase", "EN - Fase", "ES - Fase"]
            table(
                production,
                spec,
                heads,
                [[acronym, "Relações Internacionais", "Em revisão", "Em produção", "Em produção"]],
            )
        else:
            heads = ["Sigla", "Nome do arquivo", "Texto alt PT-BR"]
            table(production, spec, heads, [[acronym, f"{acronym}_00", 'Logo "São José"']])
    creation_path = tmp_path / "creation.xlsx"
    production_path = tmp_path / "production.xlsx"
    creation.save(creation_path)
    production.save(production_path)
    return creation_path, production_path
