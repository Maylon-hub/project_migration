"""XLSX adapters read named tables, never a generic first sheet."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from studybr.models import Conflict, Observation, Provenance


class SourceError(ValueError):
    pass


def clean_value(value: Any) -> Any:
    if isinstance(value, str):
        return unicodedata.normalize("NFC", value).strip() or None
    return value


@dataclass(frozen=True)
class TableSpec:
    sheet: str
    table: str
    required: tuple[str, ...]
    language: str | None = None
    role: str = ""


CREATION_TABLES = (
    TableSpec(
        "Português - home",
        "Tabela1",
        ("Sigla", "Universidade", "nome_curto", "Onde criar"),
        "pt-br",
        "home",
    ),
    TableSpec(
        "Português (subpáginas)",
        "Tabela15",
        ("Sigla", "tipo3", "nome_curto", "Onde criar"),
        "pt-br",
        "subpage",
    ),
    TableSpec(
        "Inglês (home)",
        "Tabela1346",
        ("Sigla", "Nome Inglês", "nome_curto", "Onde criar"),
        "en",
        "home",
    ),
    TableSpec(
        "Inglês (subpáginas)",
        "Tabela134",
        ("Sigla", "tipo3", "nome_curto", "Onde criar"),
        "en",
        "subpage",
    ),
    TableSpec(
        "Espanhol (home)",
        "Tabela137",
        ("Sigla", "Nome Espanhol", "nome_curto", "Onde criar"),
        "es",
        "home",
    ),
    TableSpec(
        "Espanhol (subpáginas)",
        "Tabela13",
        ("Sigla", "Tipo3", "nome_curto", "Onde criar"),
        "es",
        "subpage",
    ),
    TableSpec(
        "Compilado home",
        "Tabela7",
        ("Tipo página", "Onde criar", "Título da página", "nome_curto"),
        role="compiled_home",
    ),
    TableSpec(
        "compilado subpáginas",
        "Tabela8",
        ("tipo página", "Onde criar", "Título da página", "nome_curto"),
        role="compiled_subpage",
    ),
)
PRODUCTION_TABLES = (
    TableSpec(
        "siglas,nomes pt en es",
        "Tabela19",
        ("Siglas", "Oficial", "Inglês", "Espanhol"),
        role="names",
    ),
    TableSpec(
        "IES", "Tabela1", ("IES", "Página", "PT - Fase", "EN - Fase", "ES - Fase"), role="editorial"
    ),
    TableSpec(
        "Fotos_metadados",
        "Tabela113",
        ("Sigla", "Nome do arquivo", "Texto alt PT-BR"),
        role="media",
    ),
)


@dataclass(frozen=True)
class SourceRow:
    spec: TableSpec
    observations: dict[str, Observation]

    def value(self, key: str) -> Any:
        obs = self.observations.get(key)
        return obs.value if obs is not None else None


class WorkbookSource:
    def __init__(self, path: Path, source_id: str) -> None:
        self.path = path
        self.source_id = source_id
        self.sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        try:
            self.formulas = load_workbook(path, data_only=False, keep_links=False)
            self.cached = load_workbook(path, data_only=True, keep_links=False)
        except Exception:
            raise SourceError("Unable to read the supplied XLSX source") from None

    def rows(self, spec: TableSpec, conflicts: list[Conflict]) -> list[SourceRow]:
        if spec.sheet not in self.formulas.sheetnames:
            raise SourceError(f"Required worksheet missing: {spec.sheet}")
        sheet = self.formulas[spec.sheet]
        if spec.table not in sheet.tables:
            raise SourceError(f"Required table missing: {spec.sheet}/{spec.table}")
        min_col, min_row, max_col, max_row = range_boundaries(sheet.tables[spec.table].ref)
        heads = [clean_value(sheet.cell(min_row, c).value) for c in range(min_col, max_col + 1)]
        if any(not isinstance(h, str) for h in heads) or len(set(heads)) != len(heads):
            raise SourceError(f"Invalid table headers: {spec.sheet}/{spec.table}")
        if not set(spec.required).issubset(heads):
            raise SourceError(f"Required columns missing: {spec.sheet}/{spec.table}")
        result = []
        for number in range(min_row + 1, max_row + 1):
            observations = {}
            for col, header in zip(range(min_col, max_col + 1), heads, strict=True):
                cell = sheet.cell(number, col)
                cached = self.cached[spec.sheet].cell(number, col)
                formula = cell.value if cell.data_type == "f" else getattr(cell.value, "text", None)
                is_formula = formula is not None
                value = clean_value(cached.value if is_formula else cell.value)
                state = "not_applicable"
                if is_formula:
                    if value is None:
                        state = "missing"
                    elif cached.data_type == "e":
                        state = "error"
                    elif re.search(r"\[\d+\]|\[[^\]]+\.xlsx\]", str(formula), re.I):
                        state = "external_dependency"
                    else:
                        state = "unverified"
                provenance = Provenance(
                    self.source_id,
                    self.path.name,
                    self.sha256,
                    spec.sheet,
                    spec.table,
                    number,
                    cell.coordinate,
                    header,
                    str(formula) if is_formula else None,
                    state,
                )
                obs = Observation(value, provenance)
                observations[header] = obs
                if is_formula:
                    conflicts.append(
                        Conflict(
                            "formula_cache",
                            f"{self.source_id}:{spec.sheet}:{number}",
                            header,
                            "Formula cache is not certified; diagnostic value only",
                            (obs,),
                        )
                    )
            if any(o.value is not None for o in observations.values()):
                result.append(SourceRow(spec, observations))
        return result

    def close(self) -> None:
        self.formulas.close()
        self.cached.close()


def read_workbook(path: Path, source_id: str, specs: tuple[TableSpec, ...]) -> tuple:
    source = WorkbookSource(path, source_id)
    conflicts: list[Conflict] = []
    try:
        rows = [row for spec in specs for row in source.rows(spec, conflicts)]
        return (
            rows,
            conflicts,
            {
                "source_id": source_id,
                "filename": path.name,
                "sha256": source.sha256,
            },
        )
    finally:
        source.close()
